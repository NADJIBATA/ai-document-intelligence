"""Build test documents in memory (no binary fixtures in the repository).

The PDF writer below produces small but valid PDF 1.4 files with the standard
Helvetica fonts, which is enough to exercise page numbers, font-size based heading
detection and running headers/footers.
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from dataclasses import dataclass

import docx

PAGE_WIDTH = 612
PAGE_HEIGHT = 792
LEFT_MARGIN = 72
TOP_MARGIN = 740


@dataclass(frozen=True)
class PdfLine:
    text: str  # an empty string inserts a paragraph gap
    size: float = 11.0
    bold: bool = False


def _pdf_string(text: str) -> bytes:
    raw = text.encode("cp1252")
    return b"(" + raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)") + b")"


def _page_stream(lines: Sequence[PdfLine]) -> bytes:
    parts = [b"BT"]
    y = float(TOP_MARGIN)
    for line in lines:
        if not line.text:
            y -= 14
            continue
        font = b"/F2" if line.bold else b"/F1"
        parts.append(font + b" %.1f Tf" % line.size)
        parts.append(b"1 0 0 1 %d %.1f Tm" % (LEFT_MARGIN, y))
        parts.append(_pdf_string(line.text) + b" Tj")
        y -= line.size * 1.35
    parts.append(b"ET")
    return b"\n".join(parts)


def build_pdf(pages: Sequence[Sequence[PdfLine]], title: str | None = None) -> bytes:
    objects: list[bytes] = []

    def add(obj: bytes) -> int:
        objects.append(obj)
        return len(objects)

    catalog_id = add(b"")  # placeholder, filled once the page tree id is known
    pages_id = add(b"")
    font_regular = add(
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    )
    font_bold = add(
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>"
    )
    page_ids = []
    for lines in pages:
        stream = _page_stream(lines)
        content_id = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page_ids.append(
            add(
                b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 %d %d] "
                b"/Resources << /Font << /F1 %d 0 R /F2 %d 0 R >> >> /Contents %d 0 R >>"
                % (pages_id, PAGE_WIDTH, PAGE_HEIGHT, font_regular, font_bold, content_id)
            )
        )
    kids = b" ".join(b"%d 0 R" % page_id for page_id in page_ids)
    objects[pages_id - 1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, len(page_ids))
    objects[catalog_id - 1] = b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id
    info_id = add(b"<< /Title %s >>" % _pdf_string(title)) if title else None

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    trailer = b"<< /Size %d /Root %d 0 R" % (len(objects) + 1, catalog_id)
    if info_id:
        trailer += b" /Info %d 0 R" % info_id
    out.write(b"trailer\n" + trailer + b" >>\nstartxref\n%d\n%%%%EOF\n" % xref)
    return out.getvalue()


def simple_pdf(*page_texts: str) -> bytes:
    """One paragraph per page, body font only."""
    return build_pdf([[PdfLine(text)] for text in page_texts])


def contract_pdf() -> bytes:
    """Three pages, headings, a running header and a numbered footer."""

    def page(number: int, body: list[PdfLine]) -> list[PdfLine]:
        return [
            PdfLine("ACME Corp - Confidential", size=9),
            PdfLine(""),
            *body,
            PdfLine(""),
            PdfLine(f"Page {number} / 3", size=9),
        ]

    return build_pdf(
        [
            page(
                1,
                [
                    PdfLine("Service Agreement", size=20, bold=True),
                    PdfLine(""),
                    PdfLine("This agreement is made between ACME Corp and Globex."),
                    PdfLine("It defines the services provided by ACME Corp."),
                ],
            ),
            page(
                2,
                [
                    PdfLine("Payment Terms", size=16, bold=True),
                    PdfLine(""),
                    PdfLine("The monthly fee is 4,500 EUR, payable within 30 days."),
                ],
            ),
            page(
                3,
                [
                    PdfLine("Termination", size=16, bold=True),
                    PdfLine(""),
                    PdfLine("Either party may terminate with sixty (60) days notice."),
                ],
            ),
        ],
        title="Service Agreement",
    )


def build_docx(
    items: Sequence[tuple[str, str]], tables: Sequence[Sequence[Sequence[str]]] = ()
) -> bytes:
    """``items`` are ``(style, text)`` pairs, e.g. ``("Heading 1", "Scope")``."""
    document = docx.Document()
    document.core_properties.title = "Test document"
    document.core_properties.author = "QA"
    for style, text in items:
        document.add_paragraph(text, style=style)
    for rows in tables:
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def long_text(paragraphs: int = 12, sentences: int = 8) -> str:
    return "\n\n".join(
        " ".join(
            f"Paragraph {p} sentence {s} explains a clause of the agreement in detail."
            for s in range(sentences)
        )
        for p in range(paragraphs)
    )
